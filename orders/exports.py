from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .summary import DeliverySummary

QUANTITY_FORMAT = "#,##0.###"
MONEY_FORMAT = '#,##0.00 "₽"'
BOLD = Font(bold=True)
HEADER_FILL = PatternFill("solid", fgColor="E8EAF0")


def workbook_to_bytes(workbook: Workbook) -> bytes:
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _title(ws, summary: DeliverySummary, subtitle: str):
    ws.append([f"{summary.producer.name}: сводка на {summary.delivery_date:%d.%m.%Y}"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([subtitle])
    if summary.pending_count:
        ws.append(
            [
                f"Внимание: ещё {summary.pending_count} заказов ждут подтверждения "
                "и не вошли в сводку."
            ]
        )
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True, color="BA2121")
    ws.append([])


def _header(ws, titles):
    ws.append(titles)
    for cell in ws[ws.max_row]:
        cell.font = BOLD
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)


def _finish(ws, widths):
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def _items_sheet(ws, summary):
    ws.title = "По позициям"
    _title(
        ws,
        summary,
        f"Заказов: {summary.orders_count}, сумма {summary.total:,.2f} ₽".replace(
            ",", " "
        ),
    )
    _header(ws, ["Категория", "Позиция", "Ед.", "Количество", "Сумма, ₽"])
    ws.freeze_panes = ws.cell(row=ws.max_row + 1, column=1)
    for item in summary.items:
        ws.append([item.category, item.name, item.unit, item.quantity, item.amount])
        row = ws.max_row
        ws.cell(row=row, column=4).number_format = QUANTITY_FORMAT
        ws.cell(row=row, column=5).number_format = MONEY_FORMAT
    ws.append(["Итого", "", "", "", summary.total])
    for cell in ws[ws.max_row]:
        cell.font = BOLD
    ws.cell(row=ws.max_row, column=5).number_format = MONEY_FORMAT
    _finish(ws, [20, 38, 8, 14, 16])


def _points_sheet(ws, summary):
    ws.title = "По точкам"
    _title(ws, summary, "Состав заказов по точкам доставки")
    for point in summary.points:
        ws.append([f"{point.customer} — {point.name}"])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True, size=12)
        orders = ", ".join(f"№{pk}" for pk in point.order_ids)
        ws.append([f"{point.address}. Заказы: {orders}"])
        for comment in point.comments:
            ws.append([f"Комментарий: {comment}"])
            ws.cell(row=ws.max_row, column=1).font = Font(italic=True)
        _header(ws, ["Позиция", "Ед.", "Количество", "Сумма, ₽"])
        for line in point.lines:
            ws.append([line.name, line.unit, line.quantity, line.amount])
            ws.cell(row=ws.max_row, column=3).number_format = QUANTITY_FORMAT
            ws.cell(row=ws.max_row, column=4).number_format = MONEY_FORMAT
        ws.append(["Итого по точке", "", "", point.total])
        for cell in ws[ws.max_row]:
            cell.font = BOLD
        ws.cell(row=ws.max_row, column=4).number_format = MONEY_FORMAT
        ws.append([])
    _finish(ws, [44, 8, 14, 16])


def _matrix_sheet(ws, summary):
    ws.title = "Раскладка"
    _title(ws, summary, "Сколько каждой позиции положить в заказ каждой точки")
    _header(ws, ["Позиция", "Ед.", *[p.name for p in summary.points], "Всего"])
    ws.freeze_panes = ws.cell(row=ws.max_row + 1, column=3)
    by_point = {
        point.point_id: {
            (line.product_id, line.unit): line.quantity for line in point.lines
        }
        for point in summary.points
    }
    for item in summary.items:
        quantities = [
            by_point[p.point_id].get((item.product_id, item.unit))
            for p in summary.points
        ]
        ws.append([item.name, item.unit, *quantities, item.quantity])
        for column in range(3, 4 + len(summary.points)):
            ws.cell(row=ws.max_row, column=column).number_format = QUANTITY_FORMAT
        ws.cell(row=ws.max_row, column=3 + len(summary.points)).font = BOLD
    _finish(ws, [38, 8, *[16] * len(summary.points), 14])


def build_summary_workbook(summary: DeliverySummary) -> Workbook:
    workbook = Workbook()
    _items_sheet(workbook.active, summary)
    _points_sheet(workbook.create_sheet(), summary)
    _matrix_sheet(workbook.create_sheet(), summary)
    return workbook
