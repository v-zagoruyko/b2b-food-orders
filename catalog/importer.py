"""Импорт каталога из Excel или CSV по шаблону.

Три шага: `read_rows` читает файл в строки, `plan_import` сверяет их с каталогом
и ничего не пишет, `apply_import` записывает проверенный план.
"""

import csv
import difflib
import io
import re
import unicodedata
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation

from django.db import transaction
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from .models import Category, Product, Unit

MAX_ROWS = 2000
MAX_FILE_SIZE = 2 * 1024 * 1024
DEFAULT_CATEGORY = "Без категории"
MAX_PRICE = Decimal("99999999.99")

COLUMNS = {
    "category": "Категория",
    "name": "Название",
    "unit": "Единица",
    "step": "Шаг заказа",
    "price": "Цена",
    "description": "Описание",
    "active": "Активна",
}
REQUIRED = ("name", "unit", "price")

# Допустимые названия колонок (после нормализации).
HEADER_ALIASES = {
    "category": {"категория", "группа", "раздел"},
    "name": {"название", "наименование", "позиция", "товар"},
    "unit": {"единица", "ед", "едизм", "единицаизмерения", "единицы"},
    "step": {"шаг", "шагзаказа", "кратность"},
    "price": {"цена", "ценаза единицу", "ценазаединицу", "стоимость", "цена₽"},
    "description": {"описание", "комментарий", "состав"},
    "active": {"активна", "активен", "впродаже", "статус"},
}
UNIT_ALIASES = {
    Unit.PIECE: {"шт", "штука", "штуки", "штук", "pcs", "pc"},
    Unit.KG: {"кг", "килограмм", "килограмма", "килограммов", "kg"},
    Unit.LITER: {"л", "литр", "литра", "литров", "l"},
    Unit.PACK: {"уп", "упак", "упаковка", "упаковки", "пачка", "пач"},
}
TRUE_VALUES = {"да", "1", "true", "yes", "y", "+", "активна", "активен", "в продаже"}
FALSE_VALUES = {"нет", "0", "false", "no", "n", "-", "неактивна", "снята"}


class ImportFileError(Exception):
    """Файл нельзя прочитать как шаблон. Текст понятен сотруднику."""


def _norm(value) -> str:
    text = "" if value is None else str(value)
    return re.sub(r"[\s.,:;]+", "", text.lower().replace("ё", "е"))


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        # Числа из формул Excel приходят с «шумом» (102.00000000000001).
        value = round(value, 6)
        if value.is_integer():
            value = int(value)
    return re.sub(r"\s+", " ", str(value)).strip()


def _sheet_to_table(rows) -> list[list]:
    return [list(row) for row in rows]


def read_rows(filename: str, content: bytes) -> list[dict]:
    """Читает файл и возвращает строки как словари {поле: значение} с номером строки."""
    if len(content) > MAX_FILE_SIZE:
        raise ImportFileError("Файл больше 2 МБ.")
    name = filename.lower()
    if name.endswith((".xlsx", ".xlsm")):
        try:
            workbook = load_workbook(
                io.BytesIO(content), read_only=True, data_only=True
            )
        except Exception as exc:
            raise ImportFileError("Не удалось открыть Excel-файл.") from exc
        table = _sheet_to_table(workbook.worksheets[0].iter_rows(values_only=True))
    elif name.endswith((".csv", ".txt")):
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            try:
                text = content.decode("cp1251")
            except UnicodeDecodeError as exc:
                raise ImportFileError("Не удалось определить кодировку CSV.") from exc
        first = text.splitlines()[0] if text.strip() else ""
        delimiter = ";" if first.count(";") >= first.count(",") else ","
        table = list(csv.reader(io.StringIO(text), delimiter=delimiter))
    else:
        raise ImportFileError("Нужен файл Excel (.xlsx) или CSV.")

    header_index = next(
        (i for i, row in enumerate(table) if any(_text(cell) for cell in row)), None
    )
    if header_index is None:
        raise ImportFileError("Файл пустой.")
    columns = {}
    for position, cell in enumerate(table[header_index]):
        key = _norm(cell)
        for field_name, aliases in HEADER_ALIASES.items():
            if key in aliases and field_name not in columns:
                columns[field_name] = position
    missing = [COLUMNS[f] for f in REQUIRED if f not in columns]
    if missing:
        raise ImportFileError(
            "В файле нет обязательных колонок: " + ", ".join(missing) + ". "
            "Скачайте шаблон и заполните его."
        )

    rows = []
    for offset, row in enumerate(table[header_index + 1 :], start=header_index + 2):
        values = {
            field_name: row[position] if position < len(row) else None
            for field_name, position in columns.items()
        }
        if not any(_text(v) for v in values.values()):
            continue
        rows.append({"line": offset, **{k: _text(v) for k, v in values.items()}})
        if len(rows) > MAX_ROWS:
            raise ImportFileError(f"Слишком много строк: максимум {MAX_ROWS}.")
    if not rows:
        raise ImportFileError("В файле нет строк с позициями.")
    return [{"_columns": sorted(columns), **row} for row in rows]


def _parse_decimal(text: str, places: int, label: str) -> Decimal:
    if not text:
        raise ValueError(f"{label} не заполнена.")
    cleaned = (
        re.sub(r"(?i)(₽|руб\.?|р\.)", "", text).replace("\xa0", "").replace(" ", "")
    )
    cleaned = cleaned.replace(",", ".")
    try:
        value = Decimal(cleaned)
    except InvalidOperation:
        raise ValueError(f"{label}: «{text}» не число.") from None
    if not value.is_finite() or value != value.quantize(Decimal(1).scaleb(-places)):
        raise ValueError(f"{label}: не больше {places} знаков после запятой.")
    return value


def _parse_unit(text: str) -> str:
    if not text:
        raise ValueError("Единица не заполнена.")
    key = _norm(text)
    for unit, aliases in UNIT_ALIASES.items():
        if key in aliases or key == _norm(unit.label):
            return unit.value
    allowed = ", ".join(unit.label for unit in Unit)
    raise ValueError(f"Единица «{text}» неизвестна. Допустимо: {allowed}.")


def _parse_active(text: str) -> bool:
    key = text.strip().lower()
    if key in TRUE_VALUES:
        return True
    if key in FALSE_VALUES:
        return False
    raise ValueError(f"Активна: «{text}» не распознано, напишите «да» или «нет».")


ACTION_LABELS = {
    "create": "создать",
    "update": "изменить",
    "unchanged": "без изменений",
    "error": "ошибка",
}


@dataclass
class RowPlan:
    line: int
    name: str
    action: str  # create | update | unchanged | error
    changes: list[str] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)
    values: dict = field(default_factory=dict)
    notes: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    position: int = 0

    @property
    def action_label(self):
        return ACTION_LABELS[self.action]


@dataclass
class ImportPlan:
    rows: list[RowPlan]
    absent_count: int

    def count(self, action: str) -> int:
        return sum(1 for row in self.rows if row.action == action)

    # Для шаблона: вызвать метод с аргументом в нём нельзя.
    count_create = property(lambda self: self.count("create"))
    count_update = property(lambda self: self.count("update"))
    count_unchanged = property(lambda self: self.count("unchanged"))
    count_error = property(lambda self: self.count("error"))

    @property
    def has_errors(self) -> bool:
        return any(row.action == "error" for row in self.rows)

    @property
    def has_changes(self) -> bool:
        return bool(self.count("create") or self.count("update"))


SIMILAR_CUTOFF = 0.88


def _key(name: str) -> str:
    """Ключ сопоставления: без регистра, ё, кавычек, дефисов и знаков препинания.

    «Борщ «Домашний»», «Борщ "Домашний"» и «борщ домашний» дают один ключ,
    поэтому повторная загрузка файла не плодит дубли.
    """
    text = unicodedata.normalize("NFKC", name).casefold().replace("ё", "е")
    text = re.sub(r"[^\w\s]|_", " ", text)
    return " ".join(text.split())


def plan_import(producer, rows: list[dict]) -> ImportPlan:
    """Сверяет строки файла с каталогом производства, ничего не записывая."""
    existing = {
        _key(p.name): p
        for p in Product.objects.filter(producer=producer).select_related("category")
    }
    seen: set[str] = set()
    plans = []
    for position, raw in enumerate(rows):
        plan = RowPlan(
            line=raw["line"],
            name=raw.get("name", ""),
            action="unchanged",
            position=position,
        )
        columns = set(raw["_columns"])
        try:
            plan.values = _clean(raw, columns, plan)
        except ValueError as exc:
            plan.errors.append(str(exc))
        key = _key(plan.name)
        if key in seen:
            plan.errors.append("Позиция с таким названием уже есть выше в файле.")
        seen.add(key)
        if plan.errors:
            plan.action = "error"
        else:
            product = existing.get(key)
            _diff(plan, product)
            _annotate(plan, product, existing)
        plans.append(plan)
    present = {_key(r.name) for r in plans}
    absent = sum(1 for k, p in existing.items() if p.is_active and k not in present)
    return ImportPlan(rows=plans, absent_count=absent)


def _clean(raw, columns, plan) -> dict:
    errors = []
    values = {}
    name = raw.get("name", "")
    if not name:
        errors.append("Название не заполнено.")
    elif len(name) > 255:
        errors.append("Название длиннее 255 символов.")
    values["name"] = name
    for label, build in (
        ("unit", lambda: _parse_unit(raw["unit"])),
        ("price", lambda: _parse_decimal(raw["price"], 2, "Цена")),
    ):
        try:
            values[label] = build()
        except (ValueError, KeyError) as exc:
            errors.append(str(exc) if isinstance(exc, ValueError) else "Не заполнено.")
    if "price" in values and not (0 <= values["price"] <= MAX_PRICE):
        errors.append("Цена вне допустимых границ.")
    step = raw.get("step", "")
    if step:
        try:
            values["order_step"] = _parse_decimal(step, 3, "Шаг заказа")
            if values["order_step"] <= 0:
                errors.append("Шаг заказа должен быть больше нуля.")
        except ValueError as exc:
            errors.append(str(exc))
    else:
        values["order_step"] = None
    values["category"] = raw.get("category") or DEFAULT_CATEGORY
    values["description"] = raw.get("description", "")
    if "active" in columns and raw.get("active"):
        try:
            values["is_active"] = _parse_active(raw["active"])
        except ValueError as exc:
            errors.append(str(exc))
    if errors:
        raise ValueError(" ".join(errors))
    return values


def _annotate(plan: RowPlan, product, existing: dict) -> None:
    """Пояснения для предпросмотра: с чем сопоставлено и на что похоже."""
    if product is not None:
        if product.name != plan.name:
            plan.notes.append(f"сопоставлено с «{product.name}»")
        return
    close = difflib.get_close_matches(
        _key(plan.name), list(existing), n=1, cutoff=SIMILAR_CUTOFF
    )
    if close:
        plan.warnings.append(
            f"похоже на существующую «{existing[close[0]].name}»: если это та же "
            "позиция, исправьте название в файле, иначе создастся дубль"
        )


FIELD_LABELS = {
    "category": "категория",
    "unit": "единица",
    "order_step": "шаг",
    "price": "цена",
    "description": "описание",
    "is_active": "активна",
}


def _display(field_name, value):
    if field_name == "unit":
        return Unit(value).label
    if field_name == "is_active":
        return "да" if value else "нет"
    if value in (None, ""):
        return "—"
    if isinstance(value, Decimal):
        return f"{value.normalize():f}"
    return str(value)


def _diff(plan: RowPlan, product) -> None:
    values = plan.values
    if product is None:
        plan.action = "create"
        plan.changes.append(
            f"цена {_display('price', values['price'])} ₽, {_display('unit', values['unit'])}"
        )
        return
    current = {
        "category": product.category.name,
        "unit": product.unit,
        "order_step": product.order_step,
        "price": product.price,
        "description": product.description,
        "is_active": product.is_active,
    }
    for field_name, label in FIELD_LABELS.items():
        if field_name not in values:
            continue
        if field_name == "description" and not values["description"]:
            continue  # пустое описание в файле не стирает существующее
        if values[field_name] != current[field_name]:
            plan.changes.append(
                f"{label}: {_display(field_name, current[field_name])} → "
                f"{_display(field_name, values[field_name])}"
            )
    plan.action = "update" if plan.changes else "unchanged"
    plan.values["_product_id"] = product.pk


@transaction.atomic
def apply_import(producer, plan: ImportPlan) -> dict:
    """Записывает проверенный план. Возвращает счётчики для сообщения."""
    if plan.has_errors:
        raise ValueError("План содержит ошибки.")
    categories = {_key(c.name): c for c in Category.objects.filter(producer=producer)}
    next_order = max((c.sort_order for c in categories.values()), default=-1) + 1

    def category_for(name):
        nonlocal next_order
        category = categories.get(_key(name))
        if category is None:
            category = Category.objects.create(
                producer=producer, name=name, sort_order=next_order
            )
            categories[_key(name)] = category
            next_order += 1
        return category

    created = updated = 0
    for row in plan.rows:
        if row.action not in ("create", "update"):
            continue
        values = row.values
        fields = {
            "category": category_for(values["category"]),
            "unit": values["unit"],
            "order_step": values["order_step"],
            "price": values["price"],
        }
        if values["description"]:
            fields["description"] = values["description"]
        if "is_active" in values:
            fields["is_active"] = values["is_active"]
        if row.action == "create":
            Product.objects.create(
                producer=producer,
                name=values["name"],
                sort_order=row.position,
                **fields,
            )
            created += 1
        else:
            product = Product.objects.get(pk=values["_product_id"], producer=producer)
            for key, value in fields.items():
                setattr(product, key, value)
            product.save()
            updated += 1
    return {"created": created, "updated": updated}


def build_template() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Каталог"
    sheet.append(list(COLUMNS.values()))
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for example in (
        ["Выпечка", "Круассан классический", "шт", "", 85, "Сливочное масло", "да"],
        ["Салаты", "Оливье", "кг", 0.5, 520, "", "да"],
        ["Напитки", "Морс клюквенный", "л", 1, 190, "", "да"],
    ):
        sheet.append(example)
    for letter, width in zip("ABCDEFG", (18, 36, 10, 12, 10, 32, 10), strict=True):
        sheet.column_dimensions[letter].width = width
    sheet.freeze_panes = "A2"

    help_sheet = workbook.create_sheet("Как заполнять")
    for line in (
        "Обязательные колонки: Название, Единица, Цена.",
        "Единица: " + ", ".join(unit.label for unit in Unit) + ".",
        "Шаг заказа: кратность количества, например 0,5 кг или 10 шт. Можно оставить пустым.",
        "Категория: если пусто, позиция попадёт в «" + DEFAULT_CATEGORY + "».",
        "Активна: да или нет. Если колонки нет или ячейка пуста, статус позиции не меняется.",
        "Позиции сопоставляются по названию. Новые создаются, изменённые обновляются.",
        "Позиции, которых нет в файле, остаются в каталоге без изменений.",
        "Перед записью система покажет, что изменится, и найдёт ошибки.",
    ):
        help_sheet.append([line])
    help_sheet.column_dimensions["A"].width = 100
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
