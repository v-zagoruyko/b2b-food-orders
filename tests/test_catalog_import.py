import io
from decimal import Decimal

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from openpyxl import Workbook, load_workbook

from catalog.importer import (
    ImportFileError,
    apply_import,
    build_template,
    plan_import,
    read_rows,
)
from catalog.models import Category, Product

from .factories import CategoryFactory, ProducerFactory, ProductFactory, make_staff

HEADER = [
    "Категория",
    "Название",
    "Единица",
    "Шаг заказа",
    "Цена",
    "Описание",
    "Активна",
]


def xlsx(rows, header=HEADER):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(header)
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def csv_bytes(text, encoding="utf-8-sig"):
    return text.encode(encoding)


@pytest.fixture
def producer(db):
    return ProducerFactory()


class TestReadRows:
    def test_reads_excel(self):
        rows = read_rows(
            "a.xlsx", xlsx([["Выпечка", "Круассан", "шт", None, 85, "", "да"]])
        )

        assert rows[0]["name"] == "Круассан"
        assert rows[0]["price"] == "85"
        assert rows[0]["line"] == 2

    def test_reads_semicolon_csv_with_decimal_comma_and_synonyms(self):
        data = csv_bytes(
            "Группа;Наименование;Ед.;Кратность;Стоимость\nСалаты;Оливье;кг;0,5;520,50\n"
        )

        row = read_rows("a.csv", data)[0]

        assert (
            row["category"],
            row["name"],
            row["unit"],
            row["step"],
            row["price"],
        ) == (
            "Салаты",
            "Оливье",
            "кг",
            "0,5",
            "520,50",
        )

    def test_reads_cp1251_csv(self):
        data = "Название;Единица;Цена\nМорс;л;190\n".encode("cp1251")

        assert read_rows("a.csv", data)[0]["name"] == "Морс"

    def test_skips_empty_rows(self):
        rows = read_rows(
            "a.xlsx",
            xlsx([["", "", "", "", "", "", ""], ["К", "А", "шт", None, 1, "", ""]]),
        )

        assert len(rows) == 1

    @pytest.mark.parametrize(
        ("filename", "data"),
        [
            ("a.pdf", b"x"),
            ("a.xlsx", b"not an excel file"),
            ("a.csv", b""),
            ("a.csv", csv_bytes("Название;Цена\nА;1\n")),  # нет колонки «Единица»
            ("a.csv", csv_bytes("Название;Единица;Цена\n")),  # нет строк
        ],
    )
    def test_rejects_bad_files(self, filename, data):
        with pytest.raises(ImportFileError):
            read_rows(filename, data)

    def test_rejects_too_large_file(self):
        with pytest.raises(ImportFileError):
            read_rows("a.csv", b"x" * (2 * 1024 * 1024 + 1))

    def test_template_is_readable_by_importer(self):
        workbook = load_workbook(io.BytesIO(build_template()))

        assert workbook.sheetnames == ["Каталог", "Как заполнять"]
        rows = read_rows("t.xlsx", build_template())
        assert [r["name"] for r in rows] == [
            "Круассан классический",
            "Оливье",
            "Морс клюквенный",
        ]


class TestPlan:
    def row(self, **kwargs):
        base = {
            "line": 2,
            "_columns": [
                "category",
                "name",
                "unit",
                "step",
                "price",
                "description",
                "active",
            ],
            "category": "Выпечка",
            "name": "Круассан",
            "unit": "шт",
            "step": "",
            "price": "85",
            "description": "",
            "active": "",
        }
        return base | kwargs

    def test_new_product_is_created(self, producer):
        plan = plan_import(producer, [self.row()])

        assert plan.rows[0].action == "create"
        assert not plan.has_errors and plan.has_changes

    def test_existing_product_shows_diff(self, producer):
        category = CategoryFactory(producer=producer, name="Выпечка")
        ProductFactory(
            producer=producer,
            category=category,
            name="Круассан",
            price=Decimal("80.00"),
            unit="PIECE",
        )

        plan = plan_import(
            producer, [self.row(name="КРУАССАН", price="85,50", unit="кг")]
        )

        row = plan.rows[0]
        assert row.action == "update"
        assert "цена: 80 → 85.5" in row.changes
        assert "единица: шт → кг" in row.changes

    def test_unchanged_row(self, producer):
        category = CategoryFactory(producer=producer, name="Выпечка")
        ProductFactory(
            producer=producer,
            category=category,
            name="Круассан",
            price=Decimal("85.00"),
        )

        plan = plan_import(producer, [self.row()])

        assert plan.rows[0].action == "unchanged"
        assert not plan.has_changes

    def test_other_producer_products_are_ignored(self, producer):
        other = ProducerFactory()
        ProductFactory(
            producer=other,
            category__producer=other,
            name="Круассан",
            price=Decimal("1"),
        )

        assert plan_import(producer, [self.row()]).rows[0].action == "create"

    @pytest.mark.parametrize(
        ("changes", "message"),
        [
            ({"name": ""}, "Название"),
            ({"unit": "ящик"}, "неизвестна"),
            ({"unit": ""}, "Единица"),
            ({"price": "abc"}, "не число"),
            ({"price": ""}, "Цена"),
            ({"price": "-5"}, "границ"),
            ({"price": "10.555"}, "знаков"),
            ({"step": "0"}, "больше нуля"),
            ({"step": "x"}, "не число"),
            ({"active": "может"}, "не распознано"),
        ],
    )
    def test_row_errors(self, producer, changes, message):
        plan = plan_import(producer, [self.row(**changes)])

        assert plan.rows[0].action == "error"
        assert message in " ".join(plan.rows[0].errors)
        assert plan.has_errors

    def test_duplicate_names_in_file(self, producer):
        plan = plan_import(producer, [self.row(), self.row(line=3, name="круассан")])

        assert [r.action for r in plan.rows] == ["create", "error"]

    def test_absent_products_are_counted_not_touched(self, producer):
        ProductFactory(
            producer=producer,
            category__producer=producer,
            name="Старый",
            is_active=True,
        )

        plan = plan_import(producer, [self.row()])

        assert plan.absent_count == 1

    def test_empty_description_does_not_erase_existing(self, producer):
        category = CategoryFactory(producer=producer, name="Выпечка")
        ProductFactory(
            producer=producer,
            category=category,
            name="Круассан",
            price=Decimal("85.00"),
            description="С маслом",
        )

        plan = plan_import(producer, [self.row(description="")])
        apply_import(producer, plan)

        assert Product.objects.get(name="Круассан").description == "С маслом"


class TestApply:
    def test_creates_categories_and_products_and_updates(self, producer):
        existing = ProductFactory(
            producer=producer,
            category__producer=producer,
            name="Оливье",
            price=Decimal("500.00"),
            unit="KG",
        )
        data = xlsx(
            [
                ["Выпечка", "Круассан", "шт", None, 85, "", "да"],
                ["Салаты", "Оливье", "кг", 0.5, 520, "", None],
                [None, "Без группы", "уп", None, 10, "", None],
            ]
        )

        result = apply_import(
            producer, plan_import(producer, read_rows("a.xlsx", data))
        )

        assert result == {"created": 2, "updated": 1}
        existing.refresh_from_db()
        assert (existing.price, existing.order_step, existing.category.name) == (
            Decimal("520.00"),
            Decimal("0.500"),
            "Салаты",
        )
        assert Product.objects.get(name="Без группы").category.name == "Без категории"
        assert Category.objects.filter(producer=producer, name="Выпечка").count() == 1
        assert existing.history.count() == 2  # создание и обновление

    def test_apply_refuses_plan_with_errors(self, producer):
        data = csv_bytes("Название;Единица;Цена\nА;ящик;1\n")

        with pytest.raises(ValueError):
            apply_import(producer, plan_import(producer, read_rows("a.csv", data)))

    def test_reimport_is_idempotent(self, producer):
        data = xlsx([["Выпечка", "Круассан", "шт", None, 85, "", "да"]])
        apply_import(producer, plan_import(producer, read_rows("a.xlsx", data)))

        plan = plan_import(producer, read_rows("a.xlsx", data))

        assert not plan.has_changes
        assert Product.objects.filter(producer=producer).count() == 1


def upload(data, name="catalog.xlsx"):
    return SimpleUploadedFile(name, data)


class TestAdminFlow:
    def test_upload_preview_confirm(self, client, producer):
        client.force_login(make_staff(producer, "Технолог"))
        data = xlsx([["Выпечка", "Круассан", "шт", None, 85, "", "да"]])

        preview = client.post(
            reverse("admin:catalog_product_import"), {"file": upload(data)}
        )
        assert preview.status_code == 200
        assert preview.context["plan"].count("create") == 1
        assert not Product.objects.exists()  # на предпросмотре ничего не пишется

        done = client.post(reverse("admin:catalog_product_import"), {"confirm": "1"})

        assert done.status_code == 302
        assert Product.objects.get().name == "Круассан"

    def test_errors_block_confirmation(self, client, producer):
        client.force_login(make_staff(producer, "Технолог"))
        data = xlsx([["Выпечка", "Круассан", "ящик", None, 85, "", "да"]])

        preview = client.post(
            reverse("admin:catalog_product_import"), {"file": upload(data)}
        )
        done = client.post(reverse("admin:catalog_product_import"), {"confirm": "1"})

        assert preview.context["plan"].has_errors
        assert done.status_code == 302
        assert not Product.objects.exists()

    def test_confirm_without_preview_does_nothing(self, client, producer):
        client.force_login(make_staff(producer, "Технолог"))

        done = client.post(reverse("admin:catalog_product_import"), {"confirm": "1"})

        assert done.status_code == 302
        assert not Product.objects.exists()

    def test_bad_file_shows_form_error(self, client, producer):
        client.force_login(make_staff(producer, "Технолог"))

        response = client.post(
            reverse("admin:catalog_product_import"), {"file": upload(b"x", "a.pdf")}
        )

        assert response.status_code == 200
        assert "Нужен файл Excel" in str(response.context["form"].errors)

    def test_imports_only_into_own_producer(self, client, producer):
        other = ProducerFactory()
        client.force_login(make_staff(producer, "Технолог"))
        data = xlsx([["Выпечка", "Круассан", "шт", None, 85, "", "да"]])

        client.post(
            reverse("admin:catalog_product_import"),
            {"file": upload(data), "producer": other.pk},
        )
        client.post(
            reverse("admin:catalog_product_import"),
            {"confirm": "1", "producer": other.pk},
        )

        assert Product.objects.get().producer == producer

    @pytest.mark.parametrize("group", ["Бухгалтер", "Менеджер заказов"])
    def test_roles_without_catalog_rights_are_denied(self, client, producer, group):
        client.force_login(make_staff(producer, group))

        assert client.get(reverse("admin:catalog_product_import")).status_code == 403
        assert (
            client.get(reverse("admin:catalog_product_import_template")).status_code
            == 403
        )

    def test_template_download(self, client, producer):
        client.force_login(make_staff(producer, "Технолог"))

        response = client.get(reverse("admin:catalog_product_import_template"))

        assert response.status_code == 200
        assert "catalog-template.xlsx" in response["Content-Disposition"]


class TestNormalization:
    def row(self, name, **kwargs):
        base = {
            "line": 2,
            "_columns": ["category", "name", "unit", "step", "price"],
            "category": "Горячее",
            "name": name,
            "unit": "шт",
            "step": "",
            "price": "140",
            "description": "",
            "active": "",
        }
        return base | kwargs

    @pytest.mark.parametrize(
        "variant",
        [
            "Борщ «Домашний»",
            'Борщ "Домашний"',
            "борщ  домашний",
            "БОРЩ “ДОМАШНИЙ”",
            "Борщ Домашний.",
        ],
    )
    def test_quote_case_and_punctuation_variants_match_one_product(
        self, producer, variant
    ):
        category = CategoryFactory(producer=producer, name="Горячее")
        ProductFactory(
            producer=producer,
            category=category,
            name="Борщ «Домашний»",
            price=Decimal("140.00"),
            unit="PIECE",
        )

        row = plan_import(producer, [self.row(variant)]).rows[0]

        assert row.action == "unchanged"
        assert not row.warnings

    def test_hyphen_and_yo_variants_match(self, producer):
        category = CategoryFactory(producer=producer, name="Горячее")
        ProductFactory(
            producer=producer,
            category=category,
            name="Котлета по-киевски",
            price=Decimal("140.00"),
            unit="PIECE",
        )

        row = plan_import(producer, [self.row("Котлета по киевски")]).rows[0]

        assert row.action == "unchanged"
        assert row.notes == ["сопоставлено с «Котлета по-киевски»"]

    def test_match_keeps_existing_name(self, producer):
        ProductFactory(
            producer=producer,
            category__producer=producer,
            name="Борщ «Домашний»",
            price=Decimal("100.00"),
        )

        plan = plan_import(producer, [self.row('борщ "домашний"', price="150")])
        apply_import(producer, plan)

        assert Product.objects.get().name == "Борщ «Домашний»"
        assert Product.objects.get().price == Decimal("150.00")

    def test_typo_gets_duplicate_warning_but_is_not_blocked(self, producer):
        ProductFactory(
            producer=producer, category__producer=producer, name="Круассан классический"
        )

        plan = plan_import(producer, [self.row("Круасан классический")])

        row = plan.rows[0]
        assert row.action == "create"
        assert "Круассан классический" in row.warnings[0]
        assert not plan.has_errors

    def test_unrelated_name_has_no_warning(self, producer):
        ProductFactory(
            producer=producer, category__producer=producer, name="Круассан классический"
        )

        assert plan_import(producer, [self.row("Оливье")]).rows[0].warnings == []

    def test_category_matches_regardless_of_case_and_punctuation(self, producer):
        apply_import(
            producer, plan_import(producer, [self.row("А", category="Горячее блюдо")])
        )

        apply_import(
            producer, plan_import(producer, [self.row("Б", category="горячее  блюдо.")])
        )

        assert Category.objects.filter(producer=producer).count() == 1

    def test_excel_formula_noise_in_price_is_rounded(self, producer):
        rows = read_rows(
            "a.xlsx",
            xlsx([["Выпечка", "Круассан", "шт", None, 102.00000000000001, "", ""]]),
        )

        row = plan_import(producer, rows).rows[0]

        assert row.action == "create"
        assert row.values["price"] == Decimal("102")
