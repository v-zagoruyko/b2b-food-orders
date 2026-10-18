import datetime
import io
from decimal import Decimal

import pytest
from django.core import mail
from django.core.exceptions import ValidationError
from django.urls import reverse
from freezegun import freeze_time
from openpyxl import load_workbook

from orders.exports import build_summary_workbook, workbook_to_bytes
from orders.models import OrderStatus, SummaryDelivery, SummaryTrigger
from orders.notifications import scheduled_summary_date, send_due_summaries
from orders.summary import build_summary

from .factories import (
    CategoryFactory,
    CustomerFactory,
    OrderFactory,
    OrderLineFactory,
    PointFactory,
    ProducerFactory,
    ProductFactory,
    make_staff,
)

D = datetime.date(2026, 10, 20)  # вторник


@pytest.fixture
def producer(db):
    return ProducerFactory(
        summary_emails=["kitchen@example.com", "boss@example.com"],
        summary_send_time=datetime.time(14, 5),
    )


@pytest.fixture
def catalog(producer):
    bakery = CategoryFactory(producer=producer, name="Выпечка", sort_order=1)
    salads = CategoryFactory(producer=producer, name="Салаты", sort_order=2)
    return {
        "croissant": ProductFactory(
            producer=producer, category=bakery, name="Круассан", price=Decimal("85")
        ),
        "olivier": ProductFactory(
            producer=producer,
            category=salads,
            name="Оливье",
            unit="KG",
            price=Decimal("520"),
        ),
    }


def place(producer, point_name, lines, status=OrderStatus.CONFIRMED, **kwargs):
    customer = CustomerFactory(producer=producer, name=f"ИП {point_name}")
    point = PointFactory(customer=customer, name=point_name)
    return place_at(point, lines, status, **kwargs)


def place_at(point, lines, status=OrderStatus.CONFIRMED, delivery_date=D, **kwargs):
    order = OrderFactory(
        point=point, delivery_date=delivery_date, status=status, **kwargs
    )
    for product, quantity in lines:
        OrderLineFactory(order=order, product=product, quantity=Decimal(quantity))
    order.refresh_from_db()
    return order


class TestBuildSummary:
    def test_sums_items_across_points_in_category_order(self, producer, catalog):
        place(
            producer,
            "Кофейня",
            [(catalog["olivier"], "2"), (catalog["croissant"], "10")],
        )
        place(producer, "Магазин", [(catalog["croissant"], "5")])

        summary = build_summary(producer, D)

        assert [(i.name, i.quantity) for i in summary.items] == [
            ("Круассан", Decimal("15")),
            ("Оливье", Decimal("2")),
        ]
        assert summary.orders_count == 2
        assert summary.total == Decimal("2315.00")  # 15*85 + 2*520

    def test_only_confirmed_and_delivered_are_included(self, producer, catalog):
        place(producer, "А", [(catalog["croissant"], "1")], OrderStatus.CONFIRMED)
        place(producer, "Б", [(catalog["croissant"], "1")], OrderStatus.DELIVERED)
        for status in (OrderStatus.NEW, OrderStatus.NEEDS_REVIEW):
            place(producer, f"В {status}", [(catalog["croissant"], "100")], status)
        for status in (OrderStatus.REJECTED, OrderStatus.CANCELED):
            place(producer, f"Г {status}", [(catalog["croissant"], "100")], status)

        summary = build_summary(producer, D)

        assert summary.orders_count == 2
        assert summary.items[0].quantity == Decimal("2")
        assert summary.pending_count == 2

    def test_other_date_and_other_producer_are_excluded(self, producer, catalog):
        place(producer, "А", [(catalog["croissant"], "1")])
        other_day = place(producer, "Б", [(catalog["croissant"], "7")])
        other_day.delivery_date = D + datetime.timedelta(days=1)
        other_day.save()
        foreign = ProducerFactory()
        place(
            foreign,
            "Чужая",
            [(ProductFactory(producer=foreign, category__producer=foreign), "9")],
        )

        summary = build_summary(producer, D)

        assert summary.orders_count == 1
        assert [p.name for p in summary.points] == ["А"]

    def test_two_orders_of_one_point_are_merged_with_comments(self, producer, catalog):
        first = place(
            producer, "Кофейня", [(catalog["croissant"], "3")], comment="до 8:00"
        )
        place_at(first.point, [(catalog["croissant"], "2")], comment="без орехов")

        point = build_summary(producer, D).points[0]

        assert len(point.order_ids) == 2
        assert point.comments == ["до 8:00", "без орехов"]
        assert [(line.name, line.quantity) for line in point.lines] == [
            ("Круассан", Decimal("5"))
        ]
        assert point.total == Decimal("425.00")

    def test_empty_day(self, producer):
        summary = build_summary(producer, D)

        assert summary.is_empty
        assert summary.items == [] and summary.total == 0


class TestExcel:
    def test_workbook_has_three_sheets_with_data(self, producer, catalog):
        place(
            producer,
            "Кофейня",
            [(catalog["olivier"], "2"), (catalog["croissant"], "10")],
        )
        place(producer, "Магазин", [(catalog["croissant"], "5")], OrderStatus.NEW)

        summary = build_summary(producer, D)
        workbook = load_workbook(
            io.BytesIO(workbook_to_bytes(build_summary_workbook(summary)))
        )

        assert workbook.sheetnames == ["По позициям", "По точкам", "Раскладка"]
        items = list(workbook["По позициям"].iter_rows(values_only=True))
        assert ("Выпечка", "Круассан", "шт", 10, 850) in items
        assert any("ждёт" in str(row[0]) or "ждут" in str(row[0]) for row in items)
        matrix = list(workbook["Раскладка"].iter_rows(values_only=True))
        header = next(row for row in matrix if row[0] == "Позиция")
        assert header[2:] == ("Кофейня", "Всего")


class TestScheduling:
    def test_date_and_time_window(self, producer):
        before = datetime.datetime(
            2026, 10, 19, 14, 0, tzinfo=datetime.UTC
        )  # 17:00 МСК
        # 14:05 МСК == 11:05 UTC
        early = datetime.datetime(2026, 10, 19, 11, 0, tzinfo=datetime.UTC)
        due = datetime.datetime(2026, 10, 19, 11, 5, tzinfo=datetime.UTC)

        assert scheduled_summary_date(producer, early) is None
        assert scheduled_summary_date(producer, due) == D
        assert scheduled_summary_date(producer, before) == D

    def test_skips_non_delivery_weekday(self, producer):
        friday = datetime.datetime(
            2026, 10, 23, 12, 0, tzinfo=datetime.UTC
        )  # сводка на сб

        assert scheduled_summary_date(producer, friday) is None

    def test_disabled_without_time_or_when_inactive(self, producer):
        due = datetime.datetime(2026, 10, 19, 12, 0, tzinfo=datetime.UTC)
        producer.summary_send_time = None
        assert scheduled_summary_date(producer, due) is None


class TestSendDueSummaries:
    @freeze_time("2026-10-19 14:10:00+03:00")
    def test_sends_once_with_excel_attachment(self, producer, catalog):
        place(producer, "Кофейня", [(catalog["croissant"], "10")])
        place(producer, "Магазин", [(catalog["croissant"], "1")], OrderStatus.NEW)

        first = send_due_summaries()
        second = send_due_summaries()

        assert len(first) == 1 and second == []
        message = mail.outbox[0]
        assert message.to == ["kitchen@example.com", "boss@example.com"]
        assert "20.10.2026" in message.subject
        assert (
            "ждут подтверждения" in message.body or "ждёт подтверждения" in message.body
        )
        assert message.attachments[0][0] == "summary-2026-10-20.xlsx"
        record = SummaryDelivery.objects.get()
        assert (record.trigger, record.orders_count, record.pending_count) == (
            SummaryTrigger.AUTO,
            1,
            1,
        )

    @freeze_time("2026-10-19 14:10:00+03:00")
    def test_nothing_to_send_when_no_orders(self, producer):
        assert send_due_summaries() == []
        assert mail.outbox == []

    @freeze_time("2026-10-19 13:00:00+03:00")
    def test_not_before_send_time(self, producer, catalog):
        place(producer, "Кофейня", [(catalog["croissant"], "10")])

        assert send_due_summaries() == []

    def test_failed_send_is_logged_and_retried_after_interval(
        self, producer, catalog, monkeypatch
    ):
        from django.core.mail import EmailMessage

        place(producer, "Кофейня", [(catalog["croissant"], "10")])

        def boom(self, *args, **kwargs):
            raise OSError("smtp down")

        with freeze_time("2026-10-19 14:10:00+03:00") as frozen:
            with monkeypatch.context() as patch:
                patch.setattr(EmailMessage, "send", boom)
                failed = send_due_summaries()
                too_early = send_due_summaries()
            frozen.tick(datetime.timedelta(minutes=11))
            retried = send_due_summaries()

        assert failed[0].error == "smtp down"
        assert too_early == []
        assert retried[0].error == ""
        assert len(mail.outbox) == 1

    def test_gives_up_after_max_attempts(self, producer, catalog, monkeypatch):
        from django.core.mail import EmailMessage

        place(producer, "Кофейня", [(catalog["croissant"], "10")])
        monkeypatch.setattr(
            EmailMessage,
            "send",
            lambda self, *a, **k: (_ for _ in ()).throw(OSError("x")),
        )

        with freeze_time("2026-10-19 14:10:00+03:00") as frozen:
            results = []
            for _ in range(7):
                results += send_due_summaries()
                frozen.tick(datetime.timedelta(minutes=11))

        assert len(results) == 5

    @freeze_time("2026-10-19 14:10:00+03:00")
    def test_broken_producer_does_not_block_others(self, producer, catalog, db):
        broken = ProducerFactory(
            timezone="Mars/Olympus",
            summary_emails=["x@example.com"],
            summary_send_time=datetime.time(14, 5),
        )
        place(producer, "Кофейня", [(catalog["croissant"], "10")])
        assert broken.pk

        sent = send_due_summaries()

        assert [record.producer for record in sent] == [producer]


class TestProducerSettings:
    @pytest.mark.parametrize(
        "changes",
        [
            {"summary_emails": [], "summary_send_time": datetime.time(15, 0)},
            {"summary_emails": ["a@example.com"], "summary_send_time": None},
            {"summary_days_ahead": 2},
            {"summary_send_time": datetime.time(13, 0)},
        ],
    )
    def test_inconsistent_settings_are_rejected(self, changes):
        values = {
            "summary_emails": ["a@example.com"],
            "summary_send_time": datetime.time(15, 0),
        } | changes
        producer = ProducerFactory.build(**values)

        with pytest.raises(ValidationError):
            producer.full_clean(exclude=["name"])

    def test_early_send_is_fine_for_same_day_summary(self):
        producer = ProducerFactory.build(
            summary_emails=["a@example.com"],
            summary_send_time=datetime.time(6, 0),
            summary_days_ahead=0,
        )

        producer.full_clean(exclude=["name"])


class TestAdminViews:
    def test_owner_sees_own_summary_only(self, client, producer, catalog):
        place(producer, "Кофейня", [(catalog["croissant"], "10")])
        foreign = ProducerFactory()
        place(
            foreign,
            "Чужая точка",
            [(ProductFactory(producer=foreign, category__producer=foreign), "9")],
        )
        client.force_login(make_staff(producer, "Владелец"))

        response = client.get(
            reverse("admin:orders_order_summary"),
            {"date": "2026-10-20", "producer": foreign.pk},
        )

        assert response.status_code == 200
        content = response.content.decode()
        assert "Кофейня" in content and "Чужая точка" not in content

    def test_xlsx_download(self, client, producer, catalog):
        place(producer, "Кофейня", [(catalog["croissant"], "10")])
        client.force_login(make_staff(producer, "Бухгалтер"))

        response = client.get(
            reverse("admin:orders_order_summary_xlsx"), {"date": "2026-10-20"}
        )

        assert response.status_code == 200
        assert "summary-2026-10-20.xlsx" in response["Content-Disposition"]
        assert (
            load_workbook(io.BytesIO(response.content)).sheetnames[0] == "По позициям"
        )

    def test_technologist_has_no_access(self, client, producer):
        client.force_login(make_staff(producer, "Технолог"))

        assert client.get(reverse("admin:orders_order_summary")).status_code == 403

    def test_manager_sends_by_email_but_accountant_cannot(
        self, client, producer, catalog
    ):
        place(producer, "Кофейня", [(catalog["croissant"], "10")])
        url = reverse("admin:orders_order_summary_send")

        client.force_login(make_staff(producer, "Бухгалтер"))
        assert client.post(url, {"date": "2026-10-20"}).status_code == 403

        client.force_login(make_staff(producer, "Менеджер заказов"))
        response = client.post(url, {"date": "2026-10-20"})

        assert response.status_code == 302
        assert len(mail.outbox) == 1
        assert SummaryDelivery.objects.get().trigger == SummaryTrigger.MANUAL
